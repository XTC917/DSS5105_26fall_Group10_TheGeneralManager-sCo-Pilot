"""Rewrite Sprint 2 slides 2 and 3 to match the Sprint 1 problem pages."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

PATH = r"c:\Users\LIUSI\Desktop\NUS\DSS5105\DSS5105_26fall_Group10_TheGeneralManager-sCo-Pilot\Presentation_Group10_Sprint2.pptx"
INK = RGBColor(0x1C, 0x19, 0x17)
MUTED = RGBColor(0x5C, 0x56, 0x4E)
CARD = RGBColor(0xF6, 0xF3, 0xEE)
LINE = RGBColor(0xE4, 0xDD, 0xD2)
FONT = "Calibri"

QUESTIONS = [
    ("7:00", "Did anything go wrong yesterday?", "Daily status"),
    ("9:40", "How is the TrendCart order doing?", "Order"),
    ("11:15", "800 hoodies by the 25th. Can we take it?", "Feasibility"),
    ("14:30", "It hasn't moved in a week. Chase it.", "Delay"),
    ("17:00", "Remind me if packing is still behind.", "Reminder"),
]


def set_run(run, text, size, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = etree.SubElement(r_pr, qn(tag))
        node.set("typeface", FONT)


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=INK):
    shape = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf._txBody.find(qn("a:bodyPr")).set("anchor", "t")
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(3)
        run = p.add_run()
        set_run(run, line, size, bold, color)
    return shape


def clear(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def card(slide, l, t, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1)
    return shape


def slide_two(slide):
    add_text(slide, 0.45, 0.28, 12.3, 0.28, "1   PROBLEM", 14, True, MUTED)
    add_text(slide, 0.45, 0.58, 12.3, 0.55, "The manager asks all day, but cannot write SQL", 28, True)
    x = 0.4
    for clock, question, kind in QUESTIONS:
        card(slide, x, 1.4, 2.42, 2.55)
        add_text(slide, x + 0.12, 1.52, 2.18, 0.35, clock, 18, True)
        add_text(slide, x + 0.12, 1.95, 2.18, 1.35, question, 14, True)
        add_text(slide, x + 0.12, 3.4, 2.18, 0.35, kind, 12, color=MUTED)
        x += 2.56
    add_text(
        slide, 0.45, 4.2, 12.4, 1.15,
        "SweaterCo's general manager doesn't know SQL. From the 7:00 a.m. morning briefing to the 5:00 p.m. reminder, every question means calling the planners, the data analyst.\nFive questions. One busy day. One common problem.",
        16,
    )
    add_text(
        slide, 0.45, 5.55, 12.4, 1.4,
        "If the answer, or the question we ask back, is written like code, the GM is confused.\nNo table names, no column names, and no SQL, in the answer or in the follow-up.",
        18, True,
    )


def slide_three(slide):
    add_text(slide, 0.45, 0.28, 12.3, 0.28, "1   PROBLEM", 14, True, MUTED)
    add_text(slide, 0.45, 0.58, 12.3, 0.55, "DAs are wasting time on routine and vague problems.", 28, True)
    add_text(
        slide, 0.45, 1.3, 12.4, 0.85,
        "The result? The DA loses time to repetitive requests, while the GM is left waiting for unstable answers and unable to make timely decisions.",
        18,
    )
    card(slide, 0.45, 2.25, 6.05, 1.7)
    card(slide, 6.75, 2.25, 6.1, 1.7)
    add_text(slide, 0.65, 2.4, 5.7, 0.35, "GM — Unclear questions", 16, True)
    add_text(slide, 0.65, 2.85, 5.7, 0.85, "Complex business needs, without a precise query.", 16)
    add_text(slide, 6.95, 2.4, 5.7, 0.35, "DA — Heavy workload", 16, True)
    add_text(slide, 6.95, 2.85, 5.7, 0.85, "Repeated clarification and ad-hoc SQL interrupt deeper work.", 16)
    add_text(
        slide, 0.45, 4.15, 12.4, 1.15,
        "Not a lack of data. There is a gap between the question and the answer.\nThe GM knows the business and may not know SQL. The analyst knows the data, and is already on harder work.",
        16,
    )
    add_text(slide, 0.45, 5.4, 12.4, 0.4, "01  Wait     02  Clarify again     03  Query by hand     04  Answers differ", 18, True)
    add_text(
        slide, 0.45, 6.0, 12.4, 1.05,
        "Not a dashboard. Only metrics chosen in advance.\nA code-like answer, or a code-like follow-up, leaves the GM confused. The gap stays open.",
        16,
    )


def main():
    prs = Presentation(PATH)
    for index, builder, note in (
        (1, slide_two, "Same five questions as Sprint 1. The only addition: a code-like answer or follow-up confuses the GM."),
        (2, slide_three, "Same gap as Sprint 1: the DA's time, the unclear question, wait then clarify then query by hand. Close on the code-like answer or follow-up."),
    ):
        slide = prs.slides[index]
        clear(slide)
        builder(slide)
        slide.notes_slide.notes_text_frame.text = note
    prs.save(PATH)
    print("updated slides 2 and 3")


if __name__ == "__main__":
    main()
